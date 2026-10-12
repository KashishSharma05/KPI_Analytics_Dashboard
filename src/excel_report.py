"""Write the monthly MIS report as an Excel workbook.

Run with:  python -m src.excel_report

The workbook is for business users who work in Excel or Google Sheets.
It is a working model, not a table of pasted numbers:

  Data sheets    the same monthly summary tables the dashboard reads
  MIS Report     KPIs and a month-by-month table calculated with Excel
                 formulas (SUMIFS). Pick a state in cell C4 and every
                 number and both charts follow it.
  State Ranking  revenue, share and rank of every state
  Top Categories revenue and items sold by product category

Because the numbers are formulas over the data sheets, anyone can click
a cell and see how it was calculated.
"""

import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

APP_DATA_DIR = "app/data"
OUTPUT_FILE = "excel/KPI_MIS_Report.xlsx"

GREEN = "0F766E"
TITLE_FONT = Font(size=16, bold=True, color="FFFFFF")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", start_color=GREEN)
LABEL_FONT = Font(bold=True, color="475569")
MONEY = '"R$" #,##0'
PERCENT = '0.0"%"'


def add_data_sheet(workbook, name, table_name, frame):
    """Put a DataFrame on its own sheet as an Excel table. Returns the last row number."""
    sheet = workbook.create_sheet(name)
    sheet.append(list(frame.columns))
    for row in frame.itertuples(index=False):
        sheet.append(list(row))
    last_row = len(frame) + 1
    table = Table(displayName=table_name, ref=f"A1:{get_column_letter(len(frame.columns))}{last_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    sheet.add_table(table)
    for position in range(1, len(frame.columns) + 1):
        sheet.column_dimensions[get_column_letter(position)].width = 22
    return last_row


def write_header_row(sheet, row, titles):
    for position, title in enumerate(titles, start=1):
        cell = sheet.cell(row=row, column=position, value=title)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_title(sheet, title, subtitle, width):
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    sheet["A1"] = title
    sheet["A1"].font, sheet["A1"].fill = TITLE_FONT, HEADER_FILL
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet["A2"] = subtitle
    sheet["A2"].font = Font(italic=True, color="64748B")


def build_report_sheet(sheet, months, states, last):
    """KPIs and a monthly table, all as SUMIFS formulas that follow the state picked in C4."""
    write_title(sheet, "Monthly MIS Report: Sales and Delivery",
                "Olist e-commerce, January 2017 to August 2018. Every number is a formula over the Data sheet.", 9)

    sheet["B4"], sheet["C4"] = "Customer state", "All"
    sheet["B4"].font = LABEL_FONT
    sheet["C4"].fill = PatternFill("solid", start_color="FEF3C7")
    sheet["D4"] = "<- choose a state from the list, or All"
    sheet["D4"].font = Font(italic=True, color="64748B")
    picker = DataValidation(type="list", formula1=f"=Lists!$A$1:$A${len(states) + 1}", allow_blank=False)
    sheet.add_data_validation(picker)
    picker.add("C4")
    # "*" matches every state, so one SUMIFS serves both "All" and a single state.
    sheet["I4"] = '=IF($C$4="All","*",$C$4)'
    sheet["H4"] = "filter used"
    sheet["H4"].font = Font(italic=True, color="94A3B8")

    def total(column):
        """Sum one Data column for the chosen state, over all months."""
        return f"SUMIFS(Data!${column}$2:${column}${last},Data!$B$2:$B${last},$I$4)"

    kpis = [
        ("Revenue", f"={total('E')}", MONEY),
        ("Orders", f"={total('C')}", "#,##0"),
        ("Average order value", f"=IFERROR({total('E')}/{total('D')},0)", '"R$" #,##0.00'),
        ("Late delivery rate", f"=IFERROR(100*{total('F')}/{total('G')},0)", PERCENT),
        ("Average delivery days", f"=IFERROR({total('H')}/{total('G')},0)", "0.0"),
        ("Average review score", f"=IFERROR({total('J')}/{total('I')},0)", "0.00"),
    ]
    for position, (label, formula, number_format) in enumerate(kpis):
        column = 2 + position
        sheet.cell(row=6, column=column, value=label).font = LABEL_FONT
        value = sheet.cell(row=7, column=column, value=formula)
        value.font, value.number_format = Font(size=14, bold=True, color=GREEN), number_format

    first = 10
    write_header_row(sheet, first - 1, ["Month", "Orders", "Delivered orders", "Revenue", "Revenue growth vs previous month",
                                        "Average order value", "Late delivery rate", "Average review score", "Running revenue"])
    sheet.row_dimensions[first - 1].height = 32
    for offset, month in enumerate(months):
        row = first + offset

        def month_total(column):
            """Sum one Data column for this row's month and the chosen state."""
            return (f"SUMIFS(Data!${column}$2:${column}${last},Data!$A$2:$A${last},$A{row},"
                    f"Data!$B$2:$B${last},$I$4)")

        sheet.cell(row=row, column=1, value=month.to_pydatetime()).number_format = "mmm yyyy"
        sheet.cell(row=row, column=2, value=f"={month_total('C')}").number_format = "#,##0"
        sheet.cell(row=row, column=3, value=f"={month_total('D')}").number_format = "#,##0"
        sheet.cell(row=row, column=4, value=f"={month_total('E')}").number_format = MONEY
        growth = "" if offset == 0 else f"=IFERROR(100*(D{row}-D{row - 1})/D{row - 1},0)"
        sheet.cell(row=row, column=5, value=growth).number_format = PERCENT
        sheet.cell(row=row, column=6, value=f"=IFERROR(D{row}/C{row},0)").number_format = '"R$" #,##0.00'
        sheet.cell(row=row, column=7, value=f"=IFERROR(100*{month_total('F')}/{month_total('G')},0)").number_format = PERCENT
        sheet.cell(row=row, column=8, value=f"=IFERROR({month_total('J')}/{month_total('I')},0)").number_format = "0.00"
        sheet.cell(row=row, column=9, value=f"=SUM($D${first}:D{row})").number_format = MONEY
    final = first + len(months) - 1

    for position in range(1, 10):
        sheet.column_dimensions[get_column_letter(position)].width = 20
    sheet.freeze_panes = f"A{first}"

    revenue_chart = BarChart()
    revenue_chart.title, revenue_chart.height, revenue_chart.width = "Revenue by month", 8, 18
    revenue_chart.add_data(Reference(sheet, min_col=4, min_row=first - 1, max_row=final), titles_from_data=True)
    revenue_chart.set_categories(Reference(sheet, min_col=1, min_row=first, max_row=final))
    revenue_chart.legend = None
    sheet.add_chart(revenue_chart, f"A{final + 3}")

    late_chart = LineChart()
    late_chart.title, late_chart.height, late_chart.width = "Late delivery rate by month (%)", 8, 18
    late_chart.add_data(Reference(sheet, min_col=7, min_row=first - 1, max_row=final), titles_from_data=True)
    late_chart.set_categories(Reference(sheet, min_col=1, min_row=first, max_row=final))
    late_chart.legend = None
    sheet.add_chart(late_chart, f"F{final + 3}")


def build_state_sheet(sheet, states, last):
    write_title(sheet, "State Ranking", "Totals for the whole period, ranked by revenue.", 7)
    write_header_row(sheet, 4, ["State", "Orders", "Revenue", "Share of revenue", "Rank by revenue",
                                "Late delivery rate", "Average review score"])
    first, final = 5, 4 + len(states)
    for offset, state in enumerate(states):
        row = first + offset

        def state_total(column):
            return f"SUMIFS(Data!${column}$2:${column}${last},Data!$B$2:$B${last},$A{row})"

        sheet.cell(row=row, column=1, value=state)
        sheet.cell(row=row, column=2, value=f"={state_total('C')}").number_format = "#,##0"
        sheet.cell(row=row, column=3, value=f"={state_total('E')}").number_format = MONEY
        sheet.cell(row=row, column=4, value=f"=100*C{row}/SUM($C${first}:$C${final})").number_format = PERCENT
        sheet.cell(row=row, column=5, value=f"=RANK(C{row},$C${first}:$C${final})")
        sheet.cell(row=row, column=6, value=f"=IFERROR(100*{state_total('F')}/{state_total('G')},0)").number_format = PERCENT
        sheet.cell(row=row, column=7, value=f"=IFERROR({state_total('J')}/{state_total('I')},0)").number_format = "0.00"
    for position in range(1, 8):
        sheet.column_dimensions[get_column_letter(position)].width = 20
    sheet.freeze_panes = "A5"

    chart = BarChart()
    chart.type, chart.title, chart.height, chart.width = "bar", "Top 10 states by revenue", 9, 16
    chart.add_data(Reference(sheet, min_col=3, min_row=4, max_row=14), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=5, max_row=14))
    chart.legend = None
    sheet.add_chart(chart, "I4")


def build_category_sheet(sheet, categories, last):
    write_title(sheet, "Top Categories", "Delivered orders, whole period, all states.", 5)
    write_header_row(sheet, 4, ["Category", "Orders", "Items sold", "Revenue", "Share of revenue"])
    first, final = 5, 4 + len(categories)
    for offset, category in enumerate(categories):
        row = first + offset

        def category_total(column):
            return f"SUMIFS(CategoryData!${column}$2:${column}${last},CategoryData!$B$2:$B${last},$A{row})"

        sheet.cell(row=row, column=1, value=category)
        sheet.cell(row=row, column=2, value=f"={category_total('C')}").number_format = "#,##0"
        sheet.cell(row=row, column=3, value=f"={category_total('D')}").number_format = "#,##0"
        sheet.cell(row=row, column=4, value=f"={category_total('E')}").number_format = MONEY
        sheet.cell(row=row, column=5, value=f"=100*D{row}/SUM($D${first}:$D${final})").number_format = PERCENT
    sheet.column_dimensions["A"].width = 38
    for position in range(2, 6):
        sheet.column_dimensions[get_column_letter(position)].width = 18
    sheet.freeze_panes = "A5"


def main():
    by_state = pd.read_csv(f"{APP_DATA_DIR}/monthly_state.csv", parse_dates=["month_start"])
    by_category = pd.read_csv(f"{APP_DATA_DIR}/monthly_category_state.csv", parse_dates=["month_start"])
    # One row per month and category is enough for this report, and keeps the file small.
    by_category = by_category.groupby(["month_start", "category"], as_index=False)[["orders", "items_sold", "revenue"]].sum()

    months = sorted(by_state["month_start"].unique())
    months = [pd.Timestamp(month) for month in months]
    # States and categories are listed biggest first, so the rankings read top-down.
    states = by_state.groupby("customer_state")["revenue"].sum().sort_values(ascending=False).index.tolist()
    categories = by_category.groupby("category")["revenue"].sum().sort_values(ascending=False).index.tolist()

    workbook = Workbook()
    report = workbook.active
    report.title = "MIS Report"
    state_sheet = workbook.create_sheet("State Ranking")
    category_sheet = workbook.create_sheet("Top Categories")
    last_data_row = add_data_sheet(workbook, "Data", "MonthlyStateData", by_state)
    last_category_row = add_data_sheet(workbook, "CategoryData", "MonthlyCategoryData", by_category)

    lists = workbook.create_sheet("Lists")  # the choices offered in the state picker
    for position, value in enumerate(["All"] + sorted(states), start=1):
        lists.cell(row=position, column=1, value=value)
    lists.sheet_state = "hidden"

    build_report_sheet(report, months, states, last_data_row)
    build_state_sheet(state_sheet, states, last_data_row)
    build_category_sheet(category_sheet, categories, last_category_row)

    os.makedirs(os.path.dirname(OUTPUT_FILE), exist_ok=True)
    workbook.save(OUTPUT_FILE)
    print(f"Wrote {OUTPUT_FILE}: {len(months)} months, {len(states)} states, {len(categories)} categories "
          f"({os.path.getsize(OUTPUT_FILE) / 1e3:.0f} KB)")


if __name__ == "__main__":
    main()
